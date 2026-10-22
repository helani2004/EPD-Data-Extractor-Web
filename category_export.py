"""Export Combined plus all 63 material category sheets."""

import re
from copy import copy
from datetime import datetime, timedelta, timezone

from openpyxl.utils import get_column_letter
from record_fields import excel_values
from openpyxl import load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from extractor_core import (
    TEMPLATE,
    HEADERS,
    GROUPS,
    STARTS,
    IMPACT_FIELDS,
)
from material_categories import CATEGORIES


def category_sheet_name(number, category):
    """Create a unique worksheet name within Excel's limits."""
    cleaned = re.sub(r"[\[\]:*?/\\]", "-", category)
    cleaned = cleaned.strip().strip("'")

    # The number uses three characters, leaving 28 for the name.
    return f"{number:02d} {cleaned[:28]}"


def append_values(sheet, row):
    """Write extracted text as values, never Excel formulas."""
    cleaned_row = [
        ILLEGAL_CHARACTERS_RE.sub("", value)
        if isinstance(value, str)
        else value
        for value in row
    ]

    sheet.append(cleaned_row)

    for cell, value in zip(sheet[sheet.max_row], cleaned_row):
        if isinstance(value, str):
            cell.data_type = "s"

def session_started(record):
    """Return the session start as an Excel date in Sri Lankan time."""
    raw = record['result'].get('_visit', {}).get('started_at')

    # Older records have no recorded session start.
    if not raw:
        return None

    try:
        started = datetime.fromisoformat(
            raw.replace('Z', '+00:00')
        )

        if started.tzinfo is None:
            raise ValueError('Missing timezone')

    except (TypeError, ValueError) as exc:
        raise ValueError(
            'A record has an invalid session start time.'
        ) from exc

    sri_lanka = timezone(timedelta(hours=5, minutes=30))

    # Excel stores dates without timezone information.
    return started.astimezone(sri_lanka).replace(tzinfo=None)


def add_session_column(sheet):
    """Insert the timestamp before url, preserving header layout."""
    dimensions = [
        copy(d) for d in sheet.column_dimensions.values()
    ]
    merged = [
        copy(r) for r in sheet.merged_cells.ranges
    ]

    for merged_range in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(merged_range))

    sheet.insert_cols(1)
    sheet.column_dimensions.clear()

    for dimension in dimensions:
        dimension.min += 1
        dimension.max += 1
        dimension.index = get_column_letter(dimension.min)
        sheet.column_dimensions[dimension.index] = dimension

    for merged_range in merged:
        merged_range.shift(col_shift=1)
        sheet.merge_cells(str(merged_range))

    sheet.column_dimensions['A'].width = 29

    for row in (1, 2):
        sheet.cell(row, 1)._style = copy(
            sheet.cell(row, 2)._style
        )

    sheet['A1'] = 'Session'
    sheet['A2'] = 'Session started (UTC+05:30)'
def save_category_excel(records, destination):
    """Create exactly 64 sheets using the existing Excel template."""
    if len(CATEGORIES) != 63 or len(set(CATEGORIES)) != 63:
        raise ValueError("The category list must contain 63 unique items.")

    # Check categories before creating the workbook.
    for record in records:
        category = record["result"]["data"].get(
            "Material Category", ""
        )
        if category not in CATEGORIES:
            raise ValueError(
                "Every record must have a valid Material Category."
            )

    category_rank = {
        category: index
        for index, category in enumerate(CATEGORIES)
    }

    records = sorted(
        records,
        key=lambda record: category_rank[
            record['result']['data']['Material Category']
        ],
    )

    workbook = load_workbook(TEMPLATE)
    combined = workbook.active

    # Keep only the template's main sheet.
    for sheet in list(workbook):
        if sheet is not combined:
            workbook.remove(sheet)

    # Keep the two header rows; remove existing data.
    if combined.max_row >= 3:
        combined.delete_rows(3, combined.max_row - 2)

    combined.title = "Combined"
    add_session_column(combined)
    # Create every category sheet before adding records.
    category_sheets = {}

    for number, category in enumerate(CATEGORIES, start=1):
        sheet = workbook.copy_worksheet(combined)
        sheet.title = category_sheet_name(number, category)
        category_sheets[category] = sheet

    for record in records:
        data = record["result"]["data"]
        category = data["Material Category"]

        # Session timestamp, followed by the original 111 columns.
        row = [session_started(record)] + excel_values(data)

        append_values(combined, row)
        append_values(category_sheets[category], row)

    for sheet in workbook:
        # Keep timestamp, url and category visible when scrolling.
        sheet.freeze_panes = "D3"

        sheet.auto_filter.ref = (
            f"A2:DH{max(2, sheet.max_row)}"
        )

        for row_number in range(3, sheet.max_row + 1):
            sheet.cell(row_number, 1).number_format = (
                "dd mmm yyyy hh:mm:ss"
            )

    workbook.save(destination)
    workbook.close()
